from __future__ import annotations

import importlib.util
import unittest
from pathlib import Path


SCRIPT_PATH = Path("docx自动化生成脚本.py").resolve()


class DocxSampleGeneratorTest(unittest.TestCase):
    def test_generator_writes_expected_samples_to_docs_directory(self) -> None:
        docx = self._import_docx()
        module = self._load_generator_module()

        generated_paths = module.generate_all_samples()

        expected_names = [
            "标准正文样本.docx",
            "含图片样本.docx",
            "含超链接样本.docx",
            "含页眉页脚样本.docx",
            "样式与编号差异样本.docx",
            "settings敏感样本.docx",
            "脏文档样本.docx",
            "高风险复杂样本.docx",
        ]
        expected_paths = [module.OUTPUT_DIR / name for name in expected_names]

        self.assertEqual(generated_paths, expected_paths)
        for path in expected_paths:
            self.assertTrue(path.exists(), msg=f"missing generated file: {path}")

        image_doc = docx.Document(module.OUTPUT_DIR / "含图片样本.docx")
        image_rels = [rel for rel in image_doc.part.rels.values() if "image" in rel.reltype]
        self.assertTrue(image_rels, msg="image sample should include at least one image relationship")
        self.assertIn("图 1.", image_doc.paragraphs[-1].text)

        hyperlink_doc = docx.Document(module.OUTPUT_DIR / "含超链接样本.docx")
        hyperlink_targets = [rel.target_ref for rel in hyperlink_doc.part.rels.values() if "hyperlink" in rel.reltype]
        self.assertIn(module.HYPERLINK_URL, hyperlink_targets)

        header_footer_doc = docx.Document(module.OUTPUT_DIR / "含页眉页脚样本.docx")
        self.assertEqual(len(header_footer_doc.sections), 2)
        self.assertEqual(header_footer_doc.sections[0].header.paragraphs[0].text, "第一节 页眉占位")
        self.assertEqual(header_footer_doc.sections[1].header.paragraphs[0].text, "第二节 页眉占位")
        self.assertEqual(header_footer_doc.sections[1].footer.paragraphs[0].text, "第二节 页脚占位")

        complex_doc = docx.Document(module.OUTPUT_DIR / "高风险复杂样本.docx")
        complex_hyperlinks = [rel.target_ref for rel in complex_doc.part.rels.values() if "hyperlink" in rel.reltype]
        complex_images = [rel for rel in complex_doc.part.rels.values() if "image" in rel.reltype]
        self.assertIn(module.HYPERLINK_URL, complex_hyperlinks)
        self.assertTrue(complex_images, msg="high-risk sample should include an image relationship")

    @staticmethod
    def _import_docx():
        try:
            import docx  # type: ignore
        except Exception as exc:  # pragma: no cover
            raise unittest.SkipTest(f"python-docx not available: {exc}") from exc
        return docx

    @staticmethod
    def _load_generator_module():
        spec = importlib.util.spec_from_file_location("docx_sample_generator", SCRIPT_PATH)
        if spec is None or spec.loader is None:  # pragma: no cover
            raise AssertionError("failed to load docx sample generator module spec")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module


if __name__ == "__main__":
    unittest.main()
